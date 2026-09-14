import json
import re
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MASTER = Path("/Users/kashyap/Downloads/AI_PhD_Master_Tracker_2026-09-14.xlsx")
TREE = Path("/Users/kashyap/Downloads/PhD_Application_Tree_Tracker.xlsx")
OUT = ROOT / "data"
PUBLIC_OUT = ROOT / "public" / "data"
NOW = "2026-09-14T00:00:00.000Z"


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def slugify(*parts):
    text = " ".join(str(part or "") for part in parts)
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text.lower()).strip("-")
    return text or "application"


def rows(path, sheet_name):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    raw = list(ws.iter_rows(values_only=True))
    headers = [str(clean(value)) for value in raw[0]]
    for row in raw[1:]:
        item = {headers[index]: clean(value) for index, value in enumerate(row)}
        if any(value is not None for value in item.values()):
            yield item


def normalize_deadline(value):
    if value is None:
        return None
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    return None


def normalize_stage(source_status):
    status = (source_status or "").lower()
    if "expired" in status:
        return "expired"
    return "not-started"


def normalize_task_status(value):
    text = (value or "").strip().lower()
    return {
        "completed": "completed",
        "complete": "completed",
        "done": "completed",
        "in progress": "in-progress",
        "blocked": "blocked",
        "n/a": "not-applicable",
        "na": "not-applicable",
    }.get(text, "not-started")


def group_task(task):
    title = (task or "").lower()
    if any(word in title for word in ["reference", "referee"]):
        return "References"
    if any(word in title for word in ["submit", "submission", "proof", "confirmation", "final review"]):
        return "Submission"
    if any(word in title for word in ["cv", "letter", "proposal"]):
        return "Documents"
    return "Research & Eligibility"


def import_applications(template):
    seen = {}
    apps = []
    for row in rows(MASTER, "PhD Tracker"):
        base = slugify(row.get("Institution"), row.get("Opportunity"))
        count = seen.get(base, 0)
        seen[base] = count + 1
        app_id = base if count == 0 else f"{base}-{count + 1}"
        research = row.get("Research Area")
        priority = row.get("Priority")
        apps.append(
            {
                "id": app_id,
                "sourceRow": len(apps) + 2,
                "sourcePriority": priority,
                "priority": f"P{priority}" if isinstance(priority, int) and priority in (1, 2, 3) else None,
                "opportunity": row.get("Opportunity"),
                "institution": row.get("Institution"),
                "location": row.get("Location"),
                "researchAreas": [part.strip() for part in str(research or "").split(";") if part.strip()],
                "fitScore": row.get("Fit / 10") if isinstance(row.get("Fit / 10"), (int, float)) else None,
                "funding": row.get("Funding / stipend"),
                "deadline": normalize_deadline(row.get("Deadline")),
                "deadlineText": row.get("Deadline"),
                "sourceStatus": row.get("Status"),
                "international": row.get("International"),
                "keyNotes": row.get("Key notes"),
                "sourceVerification": row.get("Source / verification"),
                "officialUrl": None,
                "stage": normalize_stage(row.get("Status")),
                "applicationDate": None,
                "finalStatus": None,
                "notes": None,
                "tasks": [dict(task) for task in template],
                "createdAt": NOW,
                "updatedAt": NOW,
                "source": row,
            }
        )
    return apps


def import_template():
    tasks = []
    for row in rows(TREE, "02_Application_Node"):
        task_id = row.get("Task ID")
        task = row.get("Task")
        if not task_id or not task:
            continue
        tasks.append(
            {
                "id": str(task_id),
                "title": task,
                "group": group_task(task),
                "required": (str(row.get("Required?") or "").lower() == "yes"),
                "status": normalize_task_status(row.get("Done?")),
                "sourceEvidence": row.get("Source / Evidence"),
                "evidenceUrl": None,
                "notes": row.get("Notes"),
                "source": row,
            }
        )
    return tasks


def import_tree():
    nodes = []
    for row in rows(TREE, "01_Tree"):
        node_id = row.get("ID")
        title = row.get("Tree / Node")
        if not node_id or not title:
            continue
        nodes.append(
            {
                "id": str(node_id),
                "title": title,
                "type": row.get("Type"),
                "parentId": str(row.get("Parent ID")) if row.get("Parent ID") is not None else None,
                "priority": f"P{row.get('Priority')}" if row.get("Priority") is not None else None,
                "sourcePriority": row.get("Priority"),
                "status": normalize_task_status(row.get("Status")),
                "meaning": row.get("What this node means"),
                "preparation": row.get("What you need to prepare"),
                "done": row.get("Done?"),
                "notes": row.get("Notes"),
                "source": row,
            }
        )
    return nodes


def import_assets():
    assets = []
    for row in rows(TREE, "03_Core_Assets"):
        asset_id = row.get("Asset ID")
        name = row.get("Asset")
        if not asset_id or not name:
            continue
        status = normalize_task_status(row.get("Status"))
        assets.append(
            {
                "id": str(asset_id),
                "name": name,
                "format": row.get("Format"),
                "masterLocation": row.get("Master location"),
                "status": "ready" if status == "completed" else status,
                "lastUpdated": row.get("Last updated"),
                "usedFor": row.get("Used for"),
                "notes": row.get("Notes"),
                "source": row,
            }
        )
    return assets


def write(name, data):
    payload = json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    for folder in (OUT, PUBLIC_OUT):
        path = folder / name
        path.write_text(payload, encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)
    PUBLIC_OUT.mkdir(parents=True, exist_ok=True)
    template = import_template()
    applications = import_applications(template)
    tree = import_tree()
    assets = import_assets()

    write("application-template.json", template)
    write("applications.json", applications)
    write("application-tree.json", tree)
    write("core-assets.json", assets)
    summary = {
        "applications": len(applications),
        "applicationTemplateTasks": len(template),
        "treeNodes": len(tree),
        "coreAssets": len(assets),
        "applicationIdsUnique": len({app["id"] for app in applications}) == len(applications),
        "treeIdsUnique": len({node["id"] for node in tree}) == len(tree),
        "taskIdsUnique": len({task["id"] for task in template}) == len(template),
        "coreAssetIdsUnique": len({asset["id"] for asset in assets}) == len(assets),
    }
    write("import-summary.json", summary)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
